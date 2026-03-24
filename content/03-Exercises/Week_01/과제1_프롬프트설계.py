#!/usr/bin/env python3
"""과제 1: 프롬프트 설계 - DOCX 생성"""
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import sys, os

ACCENT = RGBColor(0x1B, 0x4F, 0x72)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

# === Page Setup ===
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

# === Default Font ===
style = doc.styles['Normal']
font = style.font
font.name = '맑은 고딕'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

def set_font(run, name='맑은 고딕', size=None, bold=False, color=None, italic=False, underline=False):
    run.font.name = name
    run.element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size: run.font.size = Pt(size)
    run.font.bold = bold
    if color: run.font.color.rgb = color
    run.font.italic = italic
    run.font.underline = underline
    return run

def add_paragraph(text='', alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_font(run)
    return p

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_borders(cell, color="BBBBBB", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

# ==========================================
# HEADER SECTION
# ==========================================
p = add_paragraph('경희대학교 건축공학과 대학원', WD_ALIGN_PARAGRAPH.CENTER, 0, 2)
set_font(p.runs[0], size=10, color=GRAY)

p = add_paragraph('', WD_ALIGN_PARAGRAPH.CENTER, 4, 2)
run = p.add_run('과제 1: 프롬프트 설계')
set_font(run, size=22, bold=True, color=ACCENT)

p = add_paragraph('LLM활용건축공학AI구현  |  2026학년도 1학기  |  Week 01', WD_ALIGN_PARAGRAPH.CENTER, 2, 4)
set_font(p.runs[0], size=10, color=GRAY)

# Divider line
p = add_paragraph('', space_before=4, space_after=8)
p.paragraph_format.border_bottom = True
pPr = p._p.get_or_add_pPr()
pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:color="1B4F72"/></w:pBdr>')
pPr.append(pBdr)

# ==========================================
# 과제 설명
# ==========================================
p = add_paragraph('', space_before=8, space_after=10)
r = p.add_run('자신의 ')
set_font(r, size=11)
r = p.add_run('전공/관심 분야 문서를 분석')
set_font(r, size=11, bold=True)
r = p.add_run('하기 위한 ')
set_font(r, size=11)
r = p.add_run('최적 프롬프트')
set_font(r, size=11, bold=True)
r = p.add_run('를 개발하세요.')
set_font(r, size=11)

# ==========================================
# 주제 예시 (Tip Box as Table)
# ==========================================
tip_table = doc.add_table(rows=1, cols=1)
tip_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = tip_table.cell(0, 0)
set_cell_shading(cell, "E8F8F5")
# Custom borders - thick left
tc = cell._tc
tcPr = tc.get_or_add_tcPr()
borders = parse_xml(
    f'<w:tcBorders {nsdecls("w")}>'
    f'<w:top w:val="single" w:sz="4" w:color="1ABC9C"/>'
    f'<w:bottom w:val="single" w:sz="4" w:color="1ABC9C"/>'
    f'<w:left w:val="single" w:sz="18" w:color="1ABC9C"/>'
    f'<w:right w:val="single" w:sz="4" w:color="1ABC9C"/>'
    f'</w:tcBorders>'
)
tcPr.append(borders)

# Clear default paragraph
cell.paragraphs[0].clear()
p = cell.paragraphs[0]
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(4)
run = p.add_run('💡 주제 예시')
set_font(run, size=11, bold=True, color=RGBColor(0x11, 0x7A, 0x65))

examples = [
    ("건축공학", "시방서, KDS 기준서"),
    ("경영", "재무제표, 사업계획서"),
    ("법학", "판례문, 법률 조문"),
    ("이공계", "실험 보고서, 논문 초록"),
    ("인문계", "학술 논문, 원전 텍스트"),
]
for field, example in examples:
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('•  ')
    set_font(r, size=10)
    r = p.add_run(field)
    set_font(r, size=10, bold=True)
    r = p.add_run(f': {example}')
    set_font(r, size=10)

p = cell.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(6)
p.paragraph_format.left_indent = Cm(0.5)
r = p.add_run('※ 위 예시 외에도 본인 전공/관심 분야의 문서를 자유롭게 선택하세요.')
set_font(r, size=9, italic=True, color=GRAY)

add_paragraph('', space_before=4, space_after=4)

# ==========================================
# 요구사항
# ==========================================
p = add_paragraph('', space_before=12, space_after=6)
r = p.add_run('요구사항')
set_font(r, size=15, bold=True, color=ACCENT)

requirements = [
    ('6가지 기법 중 최소 3가지 이상', ' 적용', True),
    ('Claude, GPT, Gemini', ' 세 플랫폼에서 동일 프롬프트 테스트', True),
    ('각 플랫폼별 응답 품질 비교 (정확성, 상세도, 형식)', '', False),
    ('본인의 평가 및 용도별 추천', '', False),
]
for i, (main, suffix, has_bold) in enumerate(requirements, 1):
    p = add_paragraph('', space_before=3, space_after=3)
    p.paragraph_format.left_indent = Cm(0.7)
    r = p.add_run(f'{i}.  ')
    set_font(r, size=11, bold=True, color=ACCENT)
    if has_bold:
        r = p.add_run(main)
        set_font(r, size=11, bold=True)
        r = p.add_run(suffix)
        set_font(r, size=11)
    else:
        r = p.add_run(main)
        set_font(r, size=11)

# ==========================================
# 제출 양식
# ==========================================
p = add_paragraph('', space_before=16, space_after=6)
r = p.add_run('제출 양식')
set_font(r, size=15, bold=True, color=ACCENT)

# --- 1) 프롬프트 전문 ---
p = add_paragraph('', space_before=10, space_after=6)
r = p.add_run('1) 프롬프트 전문')
set_font(r, size=12, bold=True, color=RGBColor(0x2C, 0x3E, 0x50))

# Code box
code_table = doc.add_table(rows=1, cols=1)
code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = code_table.cell(0, 0)
set_cell_shading(cell, "F5F5F5")
set_cell_borders(cell, "CCCCCC", "4")
cell.paragraphs[0].clear()
p = cell.paragraphs[0]
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(10)
p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run('[여기에 작성한 프롬프트 붙여넣기]')
set_font(r, name='Consolas', size=10, italic=True, color=GRAY)
r.element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')

add_paragraph('', space_before=2, space_after=2)

# 적용 기법 체크리스트
p = add_paragraph('', space_before=8, space_after=4)
r = p.add_run('적용한 기법 체크리스트:')
set_font(r, size=11, bold=True)

techniques = ['맥락 제공', '예시 보여주기', '출력 제약 명시', '단계별 분해', '먼저 생각하도록 요청', '역할/스타일/톤 정의']
for tech in techniques:
    p = add_paragraph('', space_before=2, space_after=2)
    p.paragraph_format.left_indent = Cm(0.7)
    r = p.add_run('☐  ')
    set_font(r, size=11)
    r = p.add_run(tech)
    set_font(r, size=11)

# --- 2) 플랫폼별 비교표 ---
p = add_paragraph('', space_before=12, space_after=6)
r = p.add_run('2) 플랫폼별 비교표')
set_font(r, size=12, bold=True, color=RGBColor(0x2C, 0x3E, 0x50))

headers = ['비교 항목', 'Claude', 'GPT', 'Gemini']
rows_data = [
    '정확성 (0~5)', '상세도 (0~5)', '형식 준수 (0~5)', '주요 장점', '주요 단점'
]

table = doc.add_table(rows=1 + len(rows_data), cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Set column widths
for row in table.rows:
    row.cells[0].width = Cm(3.8)
    row.cells[1].width = Cm(3.8)
    row.cells[2].width = Cm(3.8)
    row.cells[3].width = Cm(3.8)

# Header row
for i, header in enumerate(headers):
    cell = table.rows[0].cells[i]
    set_cell_shading(cell, "D6EAF8")
    set_cell_borders(cell, "BBBBBB", "4")
    cell.paragraphs[0].clear()
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(header)
    set_font(r, size=10, bold=True, color=ACCENT)

# Data rows
for r_idx, label in enumerate(rows_data):
    for c_idx in range(4):
        cell = table.rows[r_idx + 1].cells[c_idx]
        set_cell_borders(cell, "BBBBBB", "4")
        cell.paragraphs[0].clear()
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        if c_idx == 0:
            set_cell_shading(cell, "F8F9FA")
            r = p.add_run(label)
            set_font(r, size=10, bold=True)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run('')
            set_font(r, size=10)

# --- 3) 평가 및 추천 ---
p = add_paragraph('', space_before=12, space_after=6)
r = p.add_run('3) 평가 및 추천')
set_font(r, size=12, bold=True, color=RGBColor(0x2C, 0x3E, 0x50))

p = add_paragraph('', space_before=4, space_after=4)
p.paragraph_format.left_indent = Cm(0.5)
r = p.add_run('•  종합 평가: ')
set_font(r, size=11, bold=True)
r = p.add_run('________________________________________')
set_font(r, size=11, color=RGBColor(0xCC, 0xCC, 0xCC))

p = add_paragraph('', space_before=6, space_after=2)
p.paragraph_format.left_indent = Cm(0.5)
r = p.add_run('•  용도별 추천:')
set_font(r, size=11, bold=True)

recs = ['빠른 요약이 필요할 때 →', '정확한 기준 인용이 필요할 때 →', '상세한 분석이 필요할 때 →']
for rec in recs:
    p = add_paragraph('', space_before=2, space_after=2)
    p.paragraph_format.left_indent = Cm(1.2)
    r = p.add_run(f'•  {rec} ')
    set_font(r, size=11)
    r = p.add_run('________________')
    set_font(r, size=11, color=RGBColor(0xCC, 0xCC, 0xCC))

# ==========================================
# 제출 안내 (Red Box)
# ==========================================
p = add_paragraph('', space_before=16, space_after=6)
r = p.add_run('제출 안내')
set_font(r, size=15, bold=True, color=ACCENT)

submit_table = doc.add_table(rows=1, cols=1)
submit_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = submit_table.cell(0, 0)
set_cell_shading(cell, "FDEDEC")
tc = cell._tc
tcPr = tc.get_or_add_tcPr()
borders = parse_xml(
    f'<w:tcBorders {nsdecls("w")}>'
    f'<w:top w:val="single" w:sz="4" w:color="E74C3C"/>'
    f'<w:bottom w:val="single" w:sz="4" w:color="E74C3C"/>'
    f'<w:left w:val="single" w:sz="18" w:color="E74C3C"/>'
    f'<w:right w:val="single" w:sz="4" w:color="E74C3C"/>'
    f'</w:tcBorders>'
)
tcPr.append(borders)

cell.paragraphs[0].clear()
p = cell.paragraphs[0]
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(4)
r = p.add_run('📌 제출 방법')
set_font(r, size=11, bold=True, color=RGBColor(0xC0, 0x39, 0x2B))

p = cell.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(2)
p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run('형식:  ')
set_font(r, size=10, bold=True)
r = p.add_run('PDF 또는 Word')
set_font(r, size=10)

p = cell.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(6)
p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run('필수:  ')
set_font(r, size=10, bold=True)
r = p.add_run('각 플랫폼(Claude, GPT, Gemini) 응답 ')
set_font(r, size=10)
r = p.add_run('스크린샷 첨부')
set_font(r, size=10, bold=True, underline=True)

# ==========================================
# Footer
# ==========================================
add_paragraph('', space_before=16, space_after=2)
p = add_paragraph('', space_before=2, space_after=2)
pPr = p._p.get_or_add_pPr()
pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:color="CCCCCC"/></w:pBdr>')
pPr.append(pBdr)

p = add_paragraph('경희대학교 건축공학과  |  백장운 교수', WD_ALIGN_PARAGRAPH.CENTER, 6, 0)
set_font(p.runs[0], size=9, color=GRAY)

# Save
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "과제1_프롬프트설계.docx")
doc.save(out)
print(f"Created: {out}")
