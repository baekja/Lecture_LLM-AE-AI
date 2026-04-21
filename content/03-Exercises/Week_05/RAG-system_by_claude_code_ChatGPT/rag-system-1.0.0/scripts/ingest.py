#!/usr/bin/env python3
"""
Document Ingestion Pipeline
Converts PDF, DOCX, XLSX, PPTX to Markdown format with metadata preservation
"""

import os
import sys
import json
import hashlib
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import logging
from datetime import datetime

# Document processing libraries
import fitz  # PyMuPDF for PDF
from docx import Document as DocxDocument
import mammoth
from pptx import Presentation
import pandas as pd
import openpyxl

# Setup logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Process various document formats to Markdown"""
    
    def __init__(self, input_dir: str = "docs/originals", output_dir: str = "knowledge/md"):
        self.input_dir = Path(input_dir)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        # Track processed files to avoid reprocessing
        self.metadata_file = self.output_dir / "_metadata.json"
        self.metadata = self.load_metadata()
    
    def load_metadata(self) -> Dict:
        """Load processing metadata"""
        if self.metadata_file.exists():
            with open(self.metadata_file, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {}
    
    def save_metadata(self):
        """Save processing metadata"""
        with open(self.metadata_file, 'w', encoding='utf-8') as f:
            json.dump(self.metadata, f, indent=2, ensure_ascii=False)
    
    def get_file_hash(self, filepath: Path) -> str:
        """Calculate file hash for change detection"""
        with open(filepath, 'rb') as f:
            return hashlib.md5(f.read()).hexdigest()
    
    def process_pdf(self, filepath: Path) -> str:
        """Convert PDF to Markdown with page markers"""
        logger.info(f"Processing PDF: {filepath.name}")
        doc = fitz.open(filepath)
        markdown_content = []
        
        # Add document header
        markdown_content.append(f"# {filepath.stem}\n")
        markdown_content.append(f"**Source**: {filepath.name}\n")
        markdown_content.append(f"**Type**: PDF\n")
        markdown_content.append(f"**Pages**: {len(doc)}\n")
        markdown_content.append(f"**Processed**: {datetime.now().isoformat()}\n\n")
        markdown_content.append("---\n\n")
        
        for page_num, page in enumerate(doc, 1):
            # Page marker
            markdown_content.append(f"\n## Page {page_num}\n\n")
            
            # Extract text
            text = page.get_text()
            if text.strip():
                markdown_content.append(text)
            
            # Extract tables if present
            tables = page.find_tables()
            for table_idx, table in enumerate(tables):
                markdown_content.append(f"\n### Table {page_num}.{table_idx + 1}\n\n")
                df = pd.DataFrame(table.extract())
                markdown_content.append(df.to_markdown(index=False))
                markdown_content.append("\n")
        
        doc.close()
        return '\n'.join(markdown_content)
    
    def process_docx(self, filepath: Path) -> str:
        """Convert DOCX to Markdown"""
        logger.info(f"Processing DOCX: {filepath.name}")
        
        # Use mammoth for better formatting preservation
        with open(filepath, 'rb') as f:
            result = mammoth.convert_to_markdown(f)
            
        markdown_content = []
        markdown_content.append(f"# {filepath.stem}\n")
        markdown_content.append(f"**Source**: {filepath.name}\n")
        markdown_content.append(f"**Type**: Word Document\n")
        markdown_content.append(f"**Processed**: {datetime.now().isoformat()}\n\n")
        markdown_content.append("---\n\n")
        markdown_content.append(result.value)
        
        if result.messages:
            logger.warning(f"Conversion warnings: {result.messages}")
        
        return '\n'.join(markdown_content)
    
    def process_xlsx(self, filepath: Path) -> str:
        """Convert Excel to Markdown tables"""
        logger.info(f"Processing XLSX: {filepath.name}")
        
        markdown_content = []
        markdown_content.append(f"# {filepath.stem}\n")
        markdown_content.append(f"**Source**: {filepath.name}\n")
        markdown_content.append(f"**Type**: Excel Spreadsheet\n")
        markdown_content.append(f"**Processed**: {datetime.now().isoformat()}\n\n")
        markdown_content.append("---\n\n")
        
        # Process all sheets
        xlsx = pd.ExcelFile(filepath)
        for sheet_name in xlsx.sheet_names:
            markdown_content.append(f"\n## Sheet: {sheet_name}\n\n")
            df = pd.read_excel(filepath, sheet_name=sheet_name)
            
            # Handle empty sheets
            if df.empty:
                markdown_content.append("*(Empty sheet)*\n")
                continue
            
            # Convert to markdown table
            markdown_content.append(df.to_markdown(index=False))
            markdown_content.append("\n")
        
        return '\n'.join(markdown_content)
    
    def process_pptx(self, filepath: Path) -> str:
        """Convert PowerPoint to Markdown"""
        logger.info(f"Processing PPTX: {filepath.name}")
        
        markdown_content = []
        markdown_content.append(f"# {filepath.stem}\n")
        markdown_content.append(f"**Source**: {filepath.name}\n")
        markdown_content.append(f"**Type**: PowerPoint Presentation\n")
        markdown_content.append(f"**Processed**: {datetime.now().isoformat()}\n\n")
        markdown_content.append("---\n\n")
        
        prs = Presentation(filepath)
        
        for slide_num, slide in enumerate(prs.slides, 1):
            markdown_content.append(f"\n## Slide {slide_num}\n\n")
            
            # Extract text from shapes
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text = shape.text.strip()
                    if text:
                        # Check if it's likely a title (usually first text box)
                        if shape == slide.shapes[0] and len(text) < 100:
                            markdown_content.append(f"### {text}\n\n")
                        else:
                            markdown_content.append(f"{text}\n\n")
                
                # Extract tables
                if shape.has_table:
                    markdown_content.append("\n**Table:**\n\n")
                    table_data = []
                    for row in shape.table.rows:
                        row_data = [cell.text for cell in row.cells]
                        table_data.append(row_data)
                    
                    if table_data:
                        df = pd.DataFrame(table_data[1:], columns=table_data[0])
                        markdown_content.append(df.to_markdown(index=False))
                        markdown_content.append("\n")
        
        return '\n'.join(markdown_content)
    
    def process_py(self, filepath: Path) -> str:
        """Process Python files (already text, add metadata)"""
        logger.info(f"Processing Python: {filepath.name}")
        
        with open(filepath, 'r', encoding='utf-8') as f:
            code_content = f.read()
        
        markdown_content = []
        markdown_content.append(f"# {filepath.stem}\n")
        markdown_content.append(f"**Source**: {filepath.name}\n")
        markdown_content.append(f"**Type**: Python Code\n")
        markdown_content.append(f"**Processed**: {datetime.now().isoformat()}\n\n")
        markdown_content.append("---\n\n")
        markdown_content.append("```python\n")
        markdown_content.append(code_content)
        markdown_content.append("\n```\n")
        
        return '\n'.join(markdown_content)
    
    def process_file(self, filepath: Path) -> Optional[str]:
        """Process a single file based on its extension"""
        ext = filepath.suffix.lower()
        
        try:
            if ext == '.pdf':
                return self.process_pdf(filepath)
            elif ext in ['.docx', '.doc']:
                return self.process_docx(filepath)
            elif ext in ['.xlsx', '.xls']:
                return self.process_xlsx(filepath)
            elif ext in ['.pptx', '.ppt']:
                return self.process_pptx(filepath)
            elif ext == '.py':
                return self.process_py(filepath)
            else:
                logger.warning(f"Unsupported file type: {ext}")
                return None
        except Exception as e:
            logger.error(f"Error processing {filepath}: {e}")
            return None
    
    def process_all(self, force: bool = False):
        """Process all documents in input directory"""
        if not self.input_dir.exists():
            logger.warning(f"Input directory does not exist: {self.input_dir}")
            self.input_dir.mkdir(parents=True, exist_ok=True)
            return
        
        supported_extensions = {'.pdf', '.docx', '.doc', '.xlsx', '.xls', 
                              '.pptx', '.ppt', '.py'}
        
        files_processed = 0
        files_skipped = 0
        
        for filepath in self.input_dir.rglob('*'):
            if filepath.is_file() and filepath.suffix.lower() in supported_extensions:
                # Check if file needs processing
                file_hash = self.get_file_hash(filepath)
                file_key = str(filepath.relative_to(self.input_dir))
                
                if not force and file_key in self.metadata:
                    if self.metadata[file_key]['hash'] == file_hash:
                        logger.info(f"Skipping unchanged file: {filepath.name}")
                        files_skipped += 1
                        continue
                
                # Process file
                content = self.process_file(filepath)
                
                if content:
                    # Save markdown file
                    output_path = self.output_dir / f"{filepath.stem}.md"
                    with open(output_path, 'w', encoding='utf-8') as f:
                        f.write(content)
                    
                    # Update metadata
                    self.metadata[file_key] = {
                        'hash': file_hash,
                        'processed': datetime.now().isoformat(),
                        'output': str(output_path.relative_to(self.output_dir))
                    }
                    
                    files_processed += 1
                    logger.info(f"Saved: {output_path}")
        
        # Save metadata
        self.save_metadata()
        
        logger.info(f"Processing complete: {files_processed} processed, {files_skipped} skipped")

def main():
    """Main execution"""
    import argparse
    
    parser = argparse.ArgumentParser(description='Document Ingestion Pipeline')
    parser.add_argument('--input-dir', default='docs/originals', 
                       help='Input directory for documents')
    parser.add_argument('--output-dir', default='knowledge/md',
                       help='Output directory for Markdown files')
    parser.add_argument('--force', action='store_true',
                       help='Force reprocessing of all files')
    
    args = parser.parse_args()
    
    processor = DocumentProcessor(args.input_dir, args.output_dir)
    processor.process_all(force=args.force)

if __name__ == "__main__":
    main()